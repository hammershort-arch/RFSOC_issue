#!/usr/bin/env python3
"""Build docs/04_Presentation/PJ22045_1R8T_Rex_rev0.1.pptx (2026-10-16).
版型：同事 rev0.12 簡報 docs/reference_rev0.12/pj22045_1r8t_design.pptx（layout TITLE_DARK / CONTENT、theme、頁尾、頁碼）；
繪圖元件與配色沿用同事 source/scripts/docs/make_ppt.py。
內容來源：docs/README.md、01_System/、02_RFSoC_TX/、03_RFSoC_RX/（Rex rev0.1 文件 v0.2）。
usage: python3 make_ppt.py   （需要 python-pptx）"""
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(HERE)
TEMPLATE = os.path.join(DOCS, 'reference_rev0.12', 'pj22045_1r8t_design.pptx')
OUT = os.path.join(HERE, 'PJ22045_1R8T_Rex_rev0.1.pptx')
FIG_TX = os.path.join(DOCS, '02_RFSoC_TX', 'fig', 'tx_top.png')
FIG_RX = os.path.join(DOCS, '03_RFSoC_RX', 'fig', 'rx_top.png')
FOOT_TITLE = 'jxense.com · 16 October 2026 · source: docs/（Rex rev0.1 文件 v0.2）'
FOOT_CONTENT = 'PJ22045 1R8T · Rex rev0.1'
EA_FONT = 'Microsoft JhengHei'

NAVY = RGBColor(0x14, 0x21, 0x3D); DK2 = RGBColor(0x3A, 0x45, 0x55); LT2 = RGBColor(0xEE, 0xF1, 0xF5)
ORANGE = RGBColor(0xE0, 0x7A, 0x1F); BLUE = RGBColor(0x2B, 0x6C, 0xB0); SLATE = RGBColor(0x8F, 0xA3, 0xBF)
ORANGE_L = RGBColor(0xF0, 0xA0, 0x50); BORDER = RGBColor(0xD9, 0xDD, 0xE3); STEEL = RGBColor(0xC9, 0xD3, 0xE0)
PEACH = RGBColor(0xFD, 0xF1, 0xE6); MIST = RGBColor(0xE3, 0xE8, 0xEF); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB9, 0x1C, 0x1C); GREEN = RGBColor(0x2F, 0x85, 0x5A); PALE = RGBColor(0xDD, 0xE3, 0xEC)

prs = Presentation(TEMPLATE)
# ---- drop the template's slides, keep masters / layouts / theme -------------------------------
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
L_TITLE = next(l for l in prs.slide_layouts if l.name == 'TITLE_DARK')
L_CONTENT = next(l for l in prs.slide_layouts if l.name == 'CONTENT')
for lay, new in ((L_TITLE, FOOT_TITLE), (L_CONTENT, FOOT_CONTENT)):
    for sh in lay.shapes:
        if sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip():
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = new
            for extra in r[1:]: extra.text = ''

def _sldnum(slide, layout):
    for sh in layout.placeholders:
        if sh.placeholder_format.type is not None and 'SLIDE_NUMBER' in str(sh.placeholder_format.type):
            slide.shapes._spTree.append(copy.deepcopy(sh._element))

def _ea(run):
    """中文字型（east asian）：放在 rPr 最後（solidFill / latin 之後）。"""
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', EA_FONT)

# ---- primitives（同事 make_ppt.py，加上中文字型） ---------------------------------------------
def tb(s, x, y, w, h, text, size=15, bold=False, color=DK2, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, after=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.0); tf.margin_top = tf.margin_bottom = Inches(0.0)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if after is not None: p.space_after = Pt(after)
        segs = ln if isinstance(ln, list) else [(ln, bold, color)]
        for seg in segs:
            st, b, c = (seg + (color,))[:3] if isinstance(seg, tuple) else (seg, bold, color)
            r = p.add_run(); r.text = st; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = c
            if spacing: r.font._element.set('spc', str(spacing))
            _ea(r)
    return t

def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh

def card(s, x, y, w, h, title, body, dark=False, kicker=None, tsize=18, bsize=13, bullets=False):
    rect(s, x, y, w, h, NAVY if dark else LT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    yy = y + 0.2
    if kicker:
        tb(s, x + 0.25, yy, w - 0.5, 0.3, kicker, size=11, bold=True, color=ORANGE_L if dark else ORANGE, spacing=100); yy += 0.33
    if title:
        tb(s, x + 0.25, yy, w - 0.5, 0.5, title, size=tsize, bold=True, color=ORANGE_L if dark else NAVY); yy += 0.14 + tsize / 50
    if body:
        lines = body if isinstance(body, list) else [body]
        tb(s, x + 0.25, yy, w - 0.5, y + h - yy - 0.12, [('•  ' if bullets else '') + l for l in lines], size=bsize,
           color=PALE if dark else DK2, after=4)

def stat(s, x, y, w, h, value, label, color=BLUE, vsize=36):
    rect(s, x, y, w, h, LT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    tb(s, x + 0.25, y + 0.15, w - 0.5, 0.8, value, size=vsize, bold=True, color=color)
    tb(s, x + 0.25, y + 0.98, w - 0.5, h - 1.05, label, size=13, color=DK2)

def block(s, x, y, w, h, title, sub=None, style='navy', tsize=11, ssize=9):
    fills = {'navy': (NAVY, WHITE, STEEL), 'slate': (SLATE, NAVY, NAVY), 'orange': (ORANGE_L, NAVY, NAVY),
             'light': (LT2, NAVY, DK2), 'white': (WHITE, NAVY, DK2), 'blue': (BLUE, WHITE, MIST)}
    f, tc, sc = fills[style]
    sh = rect(s, x, y, w, h, f, line=BORDER if style in ('white', 'light') else None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.size = Pt(tsize); r.font.bold = True; r.font.color.rgb = tc; _ea(r)
    if sub:
        for ln in sub.split('\n'):
            p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = ln; r.font.size = Pt(ssize); r.font.color.rgb = sc; _ea(r)
    return sh

def arrow(s, x1, y1, x2, y2, color=DK2, width=1.25):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c

def fat_arrow(s, x, y, w, h, color):
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def num(s, x, y, n, d=0.42, color=NAVY, size=13):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = color; c.line.fill.background(); c.shadow.inherit = False
    tb(s, x, y, d, d, str(n), size=size, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def pill(s, x, y, w, h, text, fill=LT2, color=NAVY, size=11):
    rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    tb(s, x, y, w, h, text, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def _border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = tcPr.makeelement(qn(tag), {'w': '9525', 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        sf = ln.makeelement(qn('a:solidFill'), {}); c = sf.makeelement(qn('a:srgbClr'), {'val': 'D9DDE3'}); sf.append(c); ln.append(sf)
        tcPr.append(ln)

def table(s, x, y, w, rows, colw, size=12, rowh=0.36, hl=None, bold_first=False):
    nr, nc = len(rows), len(rows[0])
    shp = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for a in ('firstRow', 'bandRow'): tblPr.set(a, '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is not None: tblPr.remove(sid)
    for i, cw in enumerate(colw): tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        tbl.rows[r].height = Inches(rowh)
        for c in range(nc):
            cell = tbl.cell(r, c); cell.text = ''
            cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            run = cell.text_frame.paragraphs[0].add_run(); run.text = str(rows[r][c]); run.font.size = Pt(size)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY; run.font.color.rgb = WHITE; run.font.bold = True
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else LT2; run.font.color.rgb = DK2
                if bold_first and c == 0: run.font.bold = True; run.font.color.rgb = NAVY
                if hl and (r, c) in hl:
                    f, fc = hl[(r, c)]; cell.fill.fore_color.rgb = f; run.font.color.rgb = fc; run.font.bold = True
            _ea(run)
            _border(cell)
    return tbl

def content(kicker, title):
    s = prs.slides.add_slide(L_CONTENT)
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 100: ph.text = kicker.upper()
        elif ph.placeholder_format.idx == 101: ph.text = title
        else: continue
        for p in ph.text_frame.paragraphs:
            for r in p.runs: _ea(r)
    _sldnum(s, L_CONTENT)
    return s

def picture(s, path, x, y, w):
    return s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))

# ============================================================================================
# 1  title
s = prs.slides.add_slide(L_TITLE)
for ph in s.placeholders:
    i = ph.placeholder_format.idx
    if i == 100: ph.text = 'REX REV0.1 · 2026-10-16'
    elif i == 101: ph.text = 'PJ22045 1R8T Rex rev0.1'
    elif i == 102: ph.text = '一路 RFSoC ADC 訊號經 VU13P 廣播到 8 個 RFSoC DAC；z2 每個通道加上 Calibration 與 Gain Control'
    else: continue
    for p in ph.text_frame.paragraphs:
        for r in p.runs: _ea(r)
for i, (t, orange, w) in enumerate([('8 × cal_plus_gain_control', True, 3.3), ('NCO ±1.04 GHz', False, 2.6),
                                    ('2.3 ~ 3.3 GHz', False, 2.6), ('Timing met', False, 2.6)]):
    x = [0.6, 4.1, 6.9, 9.7][i]
    rect(s, x, 4.75, w, 0.55, ORANGE if orange else DK2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    tb(s, x, 4.75, w, 0.55, t, size=16, bold=True, color=NAVY if orange else WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 2  overview
s = content('總覽', 'rev0.1：z2 每個 DAC 前加上完整處理鏈')
for i, (v, l, c) in enumerate([('8', '個 z2 通道，各自一份 cal_plus_gain_control', BLUE), ('2.8 GHz', 'RF 中心頻率（NCO ±1.04 GHz，頻段 2.3 ~ 3.3 GHz）', BLUE),
                               ('0.000', 'ns WNS · RFSoC_TX：timing met，0 failing endpoint', GREEN), ('+0.063', 'ns WNS · RFSoC_RX：timing met', GREEN)]):
    stat(s, 0.6 + i * 3.11, 1.6, 2.81, 1.8, v, l, color=c)
table(s, 0.6, 3.7, 12.13, [
    ['板子', '來源', 'Rex rev0.1 的修改', '輸出'],
    ['RFSoC_TX（z2）', 'dev 專案 Save Project As', '8 × cal_plus_gain_control + system_reg_map；NCO ±1.04 GHz；PL + PS', 'bit / ltx / xsa / BOOT.bin'],
    ['RFSoC_RX（z1）', 'rev0.12 source/47dr1', '只改 NCO 0.84 → 1.04 GHz；補回 onebit_cdc_fifo；PL only', 'bit / ltx / xsa'],
    ['VU13P', 'rev0.12 release', '無', '沿用 rev0.12'],
], colw=[2.1, 2.45, 5.48, 2.1], size=12, rowh=0.5, bold_first=True)
tb(s, 0.6, 5.95, 12.13, 0.9, [[('全部尚未上板驗證。', True, RED), ('RFSoC_TX 本版只驗證 IF 輸出端；Gain Control 中與 TRM 相關的功能（SR、TX / RX 模式、RXAGC）待 TRM 板上線後處理。'
                                 '三塊板 NCO 必須一致，不要與 rev0.11 / rev0.12 的 z1（0.84 GHz）混用。', False, DK2)]], size=13)

# 3  system
s = content('架構', '一塊 47DR 送、VU13P 廣播、一塊 47DR 收')
def dark_card(x, y, w, h, kick, title, lines, kc=ORANGE_L):
    rect(s, x, y, w, h, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    tb(s, x + 0.25, y + 0.2, w - 0.5, 0.3, kick, size=11, bold=True, color=kc, spacing=150)
    tb(s, x + 0.25, y + 0.5, w - 0.5, 0.6, title, size=24, bold=True, color=WHITE)
    tb(s, x + 0.25, y + 1.15, w - 0.5, h - 1.25, ['•  ' + l for l in lines], size=13, color=PALE, after=4)
dark_card(0.6, 1.65, 3.3, 3.75, 'SOURCE · RFSoC_RX', 'ZU47DR · z1', ['ADC 00（m00/m01）→ 160 b + timestamp', 'Aurora X0Y1 core1（X0Y6-7）', '其他 7 個 ADC 回送同編號 DAC', 'NCO +1.04 GHz（本版修改）'])
dark_card(5.0, 1.65, 3.3, 3.75, 'RELAY · 未修改', 'VU13P', ['RX X1Y8 core1（quad 228）', 'Broadcaster 1 → 8，無 tready', 'TX X1Y4 ~ X1Y7，core 0 + 1', '沿用同事 rev0.12'], kc=SLATE)
dark_card(9.4, 1.65, 3.33, 3.75, 'OUTPUT · RFSoC_TX', 'ZU47DR · z2', ['8 個 Aurora core，依 timestamp 對齊', '8 × cal_plus_gain_control（新增）', 'Calibration + Power Fit + AGC', 'DAC s00 ~ s32，NCO −1.04 GHz'])
fat_arrow(s, 3.95, 3.45, 1.0, 0.45, ORANGE); tb(s, 3.92, 3.0, 1.1, 0.4, '1 link', size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
for k in range(3): fat_arrow(s, 8.35, 3.0 + k * 0.45, 1.0, 0.32, BLUE)
tb(s, 8.3, 2.55, 1.1, 0.4, '8 links', size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
tb(s, 0.6, 5.65, 12.13, 0.9, [[('橘：', True, ORANGE), ('z1 ADC 資料（一條 2-lane Aurora）。', False, DK2), ('藍：', True, BLUE),
    ('8 條 2-lane Aurora 送到 z2，內容相同。每條：Aurora 64B66B、Framing、2 × 25.625 Gb/s、156.25 MHz refclk。Aurora 傳輸與時間同步沿用同事設計（reference_rev0.12）。', False, DK2)]], size=13)

# 4  RFSoC_RX
s = content('RFSoC_RX · z1', 'RFSoC_RX：只改 NCO，其餘與 rev0.12 相同')
picture(s, FIG_RX, 0.6, 1.55, 12.13)
table(s, 0.6, 4.2, 7.6, [
    ['項目', 'rev0.12 47dr1', 'RFSoC_RX（rev0.1）'],
    ['RFDC ADC / DAC NCO', '0.84 / −0.84 GHz', '1.04 / −1.04 GHz（與 z2 對齊）'],
    ['onebit_cdc_fifo', 'release 清理時遺失', '補回（與 RFSoC_TX 同一份）'],
    ['design_1_wrapper', 'release 沒附 .gen', 'build 腳本重新產生並設為 top'],
    ['ip_repo', 'source/47dr1/ip_repo', 'hardware/ip_repo（專案內）'],
], colw=[2.2, 2.3, 3.1], size=12, rowh=0.44, bold_first=True, hl={(1, 2): (PEACH, ORANGE)})
card(s, 8.5, 4.2, 4.23, 2.2, 'Build 01', ['WNS +0.063 / WHS +0.010 ns', '0 critical warning', 'LUT 14.7 % / FF 10.1 % / BRAM 20.8 %（= 同事 z1）', 'PL only，沒有 PS 程式'], tsize=17, bsize=12, bullets=True)

# 5  RFSoC_TX top
s = content('RFSoC_TX · z2', 'RFSoC_TX 頂層：同事設計 + 8 個新通道')
picture(s, FIG_TX, 0.6, 1.55, 12.13)
card(s, 0.6, 4.6, 3.9, 2.15, '沿用（未修改）', ['Aurora container', 'hier_rx_x0y1 ~ 4（delay → FIFO → aligner）', 'hier_mts、timer_1', 'axi_delay、GPIO'], tsize=17, bsize=12, bullets=True)
card(s, 4.72, 4.6, 3.9, 2.15, '新增', ['8 × cal_plus_gain_control（放頂層）', '8 × rfdc_iq_combine', 'system_reg_map', 'vio_ext_inputs（暫時）'], tsize=17, bsize=12, bullets=True)
card(s, 8.83, 4.6, 3.9, 2.15, '修改', ['smartconnect 6 → 23 MI', 'ADC Outclk 30 → 240 MHz', 'NCO ±0.84 → ±1.04 GHz', '補回 onebit_cdc_fifo'], tsize=17, bsize=12, bullets=True)

# 6  channel map
s = content('RFSoC_TX · z2', '8 個通道：Aurora 或本板 ADC → DAC')
rows = [['Ch', 'in_0（Aurora）', 'in_1（ADC）', 'DAC', 'adc_clk', 'Reg map', 'Gain BRAM']]
for n, (h, c, adc, dac, ck, rb, bb) in enumerate([
        ('x0y2', 'C0', 'm00 / m01', 's00', 'clk_adc0', '0xA009_0000', '0xA009_A000'), ('x0y2', 'C1', 'm02 / m03', 's02', 'clk_adc0', '0xA009_1000', '0xA009_C000'),
        ('x0y1', 'C0', 'm10 / m11', 's10', 'clk_adc1', '0xA009_2000', '0xA009_E000'), ('x0y1', 'C1', 'm12 / m13', 's12', 'clk_adc1', '0xA009_3000', '0xA00A_0000'),
        ('x0y3', 'C0', 'm20 / m21', 's20', 'clk_adc2', '0xA009_4000', '0xA00A_2000'), ('x0y3', 'C1', 'm22 / m23', 's22', 'clk_adc2', '0xA009_5000', '0xA00A_4000'),
        ('x0y4', 'C0', 'm30 / m31', 's30', 'clk_adc3', '0xA009_6000', '0xA00A_6000'), ('x0y4', 'C1', 'm32 / m33', 's32', 'clk_adc3', '0xA009_7000', '0xA00A_8000')], 1):
    rows.append([f'ch{n}', f'hier_rx_{h} / M_AXIS_{c}', adc, dac, ck, rb, bb])
table(s, 0.6, 1.6, 12.13, rows, colw=[0.75, 3.0, 1.7, 1.0, 1.5, 2.09, 2.09], size=12, rowh=0.38, bold_first=True)
card(s, 0.6, 5.2, 5.95, 1.6, '資料格式', '160-bit AXIS：每拍 5 個複數 sample，每個 {Q[15:0], I[15:0]}；256 MHz = 1.28 GSPS。DAC 對應與同事原設計相同', tsize=17, bsize=12)
card(s, 6.78, 5.2, 5.95, 1.6, '為什麼不包 hierarchy', '同事 hier_rx 的 M_AXIS_C0/C1 是 Monitor mode；包成 hierarchy 後新 pin 也變 Monitor，Aurora 資料不會接上（validate 不報錯）', tsize=17, bsize=12)

# 7  processing chain
s = content('處理鏈', '一個通道：從輸入到 DAC')
rect(s, 0.6, 1.55, 12.13, 2.0, LT2)
tb(s, 0.7, 1.6, 6.0, 0.3, 'clk = FPGA_clk 256 MHz · 160 b', size=11, bold=True, color=BLUE)
chain = [('in_0 / in_1', 'Aurora /\n本板 ADC', 'white'), ('Source Mux', 'mem39\n0 clk', 'orange'), ('bypass', 'mem38\n0 clk', 'slate'),
         ('split', 'I / Q\n0 clk', 'slate'), ('FIR', '16-tap\n9 clk', 'navy'), ('combine', 'I/Q → 160 b\n1 clk', 'slate'),
         ('Delay', 'mem17/18\n1 + coarse', 'navy'), ('Phase', '5 × CORDIC\nmem19', 'navy'), ('Gain', 'BRAM 表\n4 clk', 'navy'),
         ('bypass mux', 'mem38\n0 clk', 'slate'), ('DAC', '×3 interp\n−1.04 GHz', 'white')]
BW, STEP = 0.92, 1.121
for i, (t, sub, st) in enumerate(chain):
    block(s, 0.6 + i * STEP, 2.05, BW, 1.05, t, sub, style=st, tsize=10, ssize=8)
    if i < len(chain) - 1: arrow(s, 0.6 + i * STEP + BW, 2.575, 0.6 + (i + 1) * STEP, 2.575)
xs = lambda i: 0.6 + i * STEP
rect(s, xs(3), 3.75, xs(9) - xs(3) + BW, 1.0, PEACH)
tb(s, xs(4) + 0.2, 3.78, 5.0, 0.3, 'adc_clk 240 MHz（agc_emu）', size=11, bold=True, color=ORANGE)
block(s, xs(3), 4.05, 1.8, 0.62, 'power_fit_top', 'pf / goal_rel', style='blue', tsize=10, ssize=8)
block(s, xs(3) + 2.1, 4.05, 2.4, 0.62, 'agc_emu', 'TXAGC · RXAGC · SR · sync', style='blue', tsize=10, ssize=8)
block(s, xs(8), 4.05, BW, 0.62, 'index mux', 'mem20 / 21', style='slate', tsize=9, ssize=8)
arrow(s, xs(3) + BW / 2, 3.1, xs(3) + BW / 2, 4.05, color=BLUE, width=1.5)
arrow(s, xs(3) + 1.8, 4.36, xs(3) + 2.1, 4.36, color=BLUE, width=1.5)
arrow(s, xs(3) + 4.5, 4.36, xs(8), 4.36, color=BLUE, width=1.5)
arrow(s, xs(8) + BW / 2, 4.05, xs(8) + BW / 2, 3.1, color=BLUE, width=1.5)
arrow(s, xs(2) + BW / 2, 3.1, xs(2) + BW / 2, 3.42, color=SLATE); arrow(s, xs(2) + BW / 2, 3.42, xs(9) + BW / 2, 3.42, color=SLATE)
arrow(s, xs(9) + BW / 2, 3.42, xs(9) + BW / 2, 3.1, color=SLATE)
tb(s, xs(9) + BW + 0.08, 3.3, 1.2, 0.3, 'mem38 = 1', size=9, color=DK2)
card(s, 0.6, 5.0, 3.9, 1.8, '開機預設', 'mem39 = 0（Aurora）、mem38 = 0（完整處理）；FIR = impulse（純延遲 7 sample）', tsize=17, bsize=12)
card(s, 4.72, 5.0, 3.9, 1.8, '固定延遲 ≈ 150 ns', '(17 + L_c + coarse) clk + (fractional + FIR group delay) sample；CORDIC L_c 待確認', tsize=17, bsize=12)
card(s, 8.83, 5.0, 3.9, 1.8, '8 通道相同', 'RTL 與 IP 設定相同；通道間差異來自 Aurora / 類比鏈路，由 Delay / Phase / Gain 補償', tsize=17, bsize=12)

# 8  source mux
s = content('Signal Source Mux', 'mem39 選來源，mem38 選要不要處理')
table(s, 0.6, 1.6, 7.6, [
    ['mem39', 'mem38', 'DAC 輸出', '用途'],
    ['0', '0', 'Aurora → FIR → Delay → Phase → Gain', '開機預設；正常運作'],
    ['0', '1', 'Aurora 原樣', '= rev0.12「Aurora 直接送 DAC」'],
    ['1', '0', '本板 ADC → FIR → Delay → Phase → Gain', '本板 loopback 量測 / 校正'],
    ['1', '1', '本板 ADC 原樣', '量 RFDC ADC → DAC 原始通道'],
], colw=[1.0, 1.0, 3.3, 2.3], size=12, rowh=0.55, hl={(1, 0): (PEACH, ORANGE), (1, 1): (PEACH, ORANGE), (1, 2): (PEACH, ORANGE), (1, 3): (PEACH, ORANGE)})
tb(s, 0.6, 4.55, 7.6, 1.2, ['mem38 = 1 時正常路徑照常運作：Power Fit / AGC / status snapshot 仍有效，可在 bypass 下觀察 pf 與 TXAGC。',
                            'XSDB：reg_ch N → reg_wr signal_source_select 1（本板 ADC）/ reg_wr bypass_axis_iq 1（直通）。'], size=13, after=6)
card(s, 8.5, 1.6, 4.23, 1.6, '純組合邏輯', '0 latency、每通道約 90 ~ 130 LUT、0 FF；沒選到的一路 tready = 1，資料丟棄', tsize=17, bsize=12)
card(s, 8.5, 3.4, 4.23, 1.6, '切換暫態', '約一個處理 latency 的期間是兩個來源混合；沒有 flush / mute', tsize=17, bsize=12)
card(s, 8.5, 5.2, 4.23, 1.6, '與 rev0.12 不同', 'rev0.12 的 Aurora 直接送 DAC；本版預設會經過處理，要恢復直通請設 mem38 = 1', tsize=17, bsize=12)

# 9  calibration
s = content('Calibration', 'FIR、Delay、Phase、Gain 四級校正')
cw = (12.13 - 3 * 0.3) / 4
for i, (k, t, b) in enumerate([
        ('mem0 ~ 16', 'In-band FIR', ['16-tap complex FIR', '係數 S0.15（0x7FFF ≈ 1.0）', '9 clk + group delay', '開機 impulse，待上板量測']),
        ('mem17 / 18', 'Delay', ['coarse 0 ~ 15 clk', 'fractional 0 ~ 4 sample（1/5 clk）', 'latency 1 + coarse clk']),
        ('mem19', 'Phase', ['5 × CORDIC Rotate', '每個 sample 旋轉 phase_offset', 'latency L_c 待確認']),
        ('BRAM 2048 × 32 b', 'Gain', ['gain index 1/16 dB → 線性 gain', 'round-half-even，飽和 ±16384', '4 clk；index = TXAGC 或 manual'])]):
    card(s, 0.6 + i * (cw + 0.3), 1.6, cw, 3.2, t, b, kicker=k, tsize=20, bsize=12, bullets=True)
table(s, 0.6, 5.1, 12.13, [
    ['ch1 資源（build 01，routed）', 'LUT', 'FF', 'RAMB36', 'DSP'],
    ['cal_plus_gain_control_ch1 合計', '14,973', '25,226', '7（+ RAMB18 × 1）', '363'],
    ['in_band_flatness（FIR）', '4,917', '10,130', '0', '320'],
    ['calibration（Delay + CORDIC + Gain + BRAM）', '6,892', '9,144', '2', '40'],
], colw=[4.93, 1.7, 1.7, 2.1, 1.7], size=12, rowh=0.4, bold_first=True, hl={(2, 4): (PEACH, ORANGE)})

# 10 power fit
s = content('Power Fit', 'Power Fit：算出把訊號搬到 goal 的增益')
card(s, 0.6, 1.6, 3.9, 2.75, 'CW（sys0 = 0）', ['每拍取 sample 2，sqrt(I² + Q²)', '區塊平均 N = 2^k（預設 256 ≈ 1.0 µs）', 'goal = min(cw_goal, −margin)，預設 −6 dBFS'], kicker='量測', tsize=19, bsize=12, bullets=True)
card(s, 0.6, 4.55, 3.9, 2.25, 'Pulse（sys0 = 1）', ['pulse_avg_amp / pulse_papr（目前由 VIO）', '不變 M 拍後自動鎖存（M = 4）', 'goal = −(margin + PAPR)'], kicker='外部輸入', tsize=19, bsize=12, bullets=True)
rect(s, 4.8, 1.6, 7.93, 1.25, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
tb(s, 5.05, 1.72, 7.4, 0.5, 'pf = round(goal − input_dBFS)', size=22, bold=True, color=WHITE)
tb(s, 5.05, 2.3, 7.4, 0.45, '1/16 dB · 飽和 ±511（±31.94 dB）· 16384 = 0 dBFS · 同時輸出 goal_rel = goal − goal_ref', size=12, color=PALE)
table(s, 4.8, 3.1, 7.93, [
    ['路徑', '延遲'],
    ['power_fit_db：鎖存 → pf / goal_rel / goal', '6 拍（23 ns）'],
    ['CW：區塊最後一個 sample → pf', '≈ 15 拍（≈ 60 ns）'],
    ['CW：更新週期（k = 8）', '256 sample ≈ 1.0 µs'],
    ['Pulse：輸入改變 → pf（M = 4）', '11 拍（43 ns）'],
    ['pf 改變 → agc_cdc 的 adc_clk 端', '≈ 25 ~ 28 ns'],
], colw=[5.33, 2.6], size=12, rowh=0.4, bold_first=True)
tb(s, 4.8, 5.65, 7.93, 1.1, [[('Power Fit 只算增益，不碰資料：', True, NAVY), ('pf 經 agc_emu → gain_index_mux → Gain_Compensation 才套用到資料。量測點是 Source Mux 選到的訊號、FIR 之前。', False, DK2)]], size=13)

# 11 TXAGC
s = content('Gain Control', 'TXAGC：類比優先，數位只補殘差')
for i, (v, l) in enumerate([('5 ~ 130 dB', '類比衰減，28 個檔位（5 dB / 4 dB 步階）'), ('0 ~ 4.94 dB', '數位補的殘差 R（有遲滯時 −H ~ 4.94 + H）'), ('1 dB', '預設換檔遲滯 H（sys16 = 0x10）')]):
    stat(s, 0.6 + i * 4.11, 1.6, 3.9, 1.65, v, l, vsize=32)
# analog state bar: state 0..17 = 130 → 45 dB (5 dB), 17..27 = 45 → 5 dB (4 dB)
bx, bw = 0.6, 12.13; seg = bw / 28
for st in range(28):
    rect(s, bx + st * seg + 0.01, 3.55, seg - 0.02, 0.42, NAVY if st < 17 else BLUE)
for st, lab in [(0, '130 dB'), (17, '45 dB'), (27, '5 dB')]:
    tb(s, bx + st * seg - (1.4 if st == 27 else 0), 4.02, 1.8, 0.3, f'state {st} · {lab}', size=10, bold=True, color=NAVY, align=PP_ALIGN.RIGHT if st == 27 else PP_ALIGN.LEFT)
rect(s, bx + 6 * seg, 3.6, 5 * seg, 0.32, NAVY); tb(s, bx + 6 * seg, 3.6, 5 * seg, 0.32, '0 ~ 17：每檔 5 dB', size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, bx + 19.5 * seg, 3.6, 5 * seg, 0.32, BLUE); tb(s, bx + 19.5 * seg, 3.6, 5 * seg, 0.32, '17 ~ 27：每檔 4 dB', size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
table(s, 0.6, 4.5, 12.13, [
    ['目標總衰減 T', '公式（1/16 dB；A + R = T，數位 gain = pf − R）'],
    ['Linear（sys4 = 0，開機預設）', '1545 + linear_offset − rx_att_index × 160 + pf − offsets + algo × 16'],
    ['Manual（sys4 = 1）', 'dac_output_ref + goal_rel + 800 − target_power × 16 + manual_offset − offsets + algo × 16'],
    ['Manual TX att（mem31 = 1）', 'agc_manual_tx_att_db × 16'],
], colw=[4.2, 7.93], size=12, rowh=0.42, bold_first=True)
tb(s, 0.6, 6.3, 12.13, 0.45, 'Linear：回波以固定總增益（標稱 0 dB）通過，pf 被完全抵銷。Manual：輸出功率對準 agc_target_power。1545 = 96.5625 dB、800 = 50 dB 是標稱值。', size=12)

# 12 modes / SR
s = content('Gain Control', '模式、shift register 與本版現況')
table(s, 0.6, 1.6, 7.4, [
    ['模式', '進入條件', 'TX 衰減欄位', '數位 gain'],
    ['RX mode', 'agc_reg5[7] = 1（開機）', 'dummy 130 dB', 'pf − R'],
    ['TX mode', 'agc_reg5[7] = 0', 'TXAGC 結果', 'pf − R，與類比換檔同步'],
    ['Bypass', 'mem28 = 1', 'mem29 state', 'pf（R = 0）'],
    ['Manual RX', 'mem36 = 1', '—', 'RX 衰減 = mem37'],
], colw=[1.4, 2.4, 1.6, 2.0], size=12, rowh=0.5, bold_first=True, hl={(1, 0): (PEACH, ORANGE)})
card(s, 0.6, 4.35, 7.4, 2.45, 'Shift register（74LVC595 × 5）', ['Reg1~4：衰減器欄位，內容一變就自動送（latest wins）', 'Reg5：TXRX_SW / Filter / PA，由 mode ctrl 控制順序與 settle',
     '數位 / 類比同步：N = 35 adc_clk（sys5）與 RCLK 同一個 edge'], tsize=17, bsize=12, bullets=True)
rect(s, 8.3, 1.6, 4.43, 5.2, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
tb(s, 8.55, 1.8, 4.0, 0.4, '本版現況（IF 驗證）', size=18, bold=True, color=ORANGE_L)
tb(s, 8.55, 2.4, 3.95, 4.2, ['•  8 通道的 Reg1~5 SER / SRCLK / RCLK 沒有接到頂層、沒有 XDC：類比換檔沒有實體效果', '•  開機在 RX mode（agc_reg5 = 0x80）、Linear',
                             '•  數位仍補 R：IF 驗證時用 mem20 = 1、mem21 固定數位 gain 避開跳檔', '•  外部輸入暫時由 vio_ext_inputs 提供', '•  TRM 上線後處理（Follow-up T1 ~ T3）'], size=12, color=PALE, after=7)

# 13 timeline
s = content('Timing', 'Pulse 輸入改變後 ≈ 290 ns 套用到資料')
rect(s, 0.6, 1.6, 3.6, 5.15, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
tb(s, 0.85, 1.85, 3.1, 0.4, '條件', size=20, bold=True, color=ORANGE_L)
tb(s, 0.85, 2.4, 3.1, 2.8, ['•  Pulse avg 與 PAPR 同拍改變（t0）', '•  M = 4（sys3）', '•  SR 週期 50 adc_clk（sys6）', '•  digital_apply_delay 35（sys5）', '•  Manual / Linear 結果相同'], size=13, color=PALE, after=7)
tb(s, 0.85, 5.55, 3.1, 1.0, 'tb_pulse_timeline 模擬值，不含 595 輸出延遲與衰減器切換時間；尚未上板量測', size=11, color=STEEL)
steps = [('Power Fit', 'pf / goal_rel / goal 輸出', '43.0 ns'), ('agc_cdc', 'adc_clk 端收到 pf / goal_rel', '67.6 ~ 71.4'),
         ('TXAGC', 'tx_att_state / 數位 gain index', '84.3 ~ 88.1'), ('SR start', 'Reg1~4 開始傳送', '88.4 ~ 92.3'),
         ('RCLK', '類比衰減鎖存', '238.4 ~ 242.2'), ('gain_index_tx', '回到 FPGA_clk → gain_index_mux', '273.4 ~ 277.3'),
         ('套用', 'Gain_Compensation 4 拍後到資料', '≈ 289 ~ 293')]
x0 = 4.55
for i, (h, d, v) in enumerate(steps):
    y = 1.65 + i * 0.73
    num(s, x0, y, i + 1, color=ORANGE if i == 6 else NAVY)
    if i < 6: arrow(s, x0 + 0.21, y + 0.42, x0 + 0.21, y + 0.73, color=SLATE, width=1)
    tb(s, x0 + 0.65, y + 0.06, 6.0, 0.4, [[(h + '   ', True, NAVY), (d, False, DK2)]], size=13)
    pill(s, 10.9, y + 0.02, 1.83, 0.4, v + (' ns' if not v.endswith('ns') else ''))

# 14 register map
s = content('Register map', 'PS 位址：原位址不變，新增從 0xA009 起')
table(s, 0.6, 1.6, 7.4, [
    ['位址', '大小', '模組'],
    ['0xA000_0000', '64 KB', 'DMA_Ctrl'], ['0xA001_0000', '64 KB', 'axi_gpio_arstn'], ['0xA002_0000', '64 KB', 'axi_gpio_pma_rst'],
    ['0xA003_0000', '64 KB', 'axi_intc_0'], ['0xA004_0000', '256 KB', 'usp_rf_data_converter_1'], ['0xA008_0000', '64 KB', 'axi_delay'],
    ['0xA009_0000 ~ 7FFF', '8 × 4 KB', 'ch1 ~ 8 addressable_reg_map'], ['0xA009_8000', '4 KB', 'system_reg_map（8 通道共用）'],
    ['0xA009_A000 ~ 0xA00A_9FFF', '8 × 8 KB', 'ch1 ~ 8 gain 表 BRAM'],
], colw=[2.8, 1.2, 3.4], size=12, rowh=0.42,
    hl={(r, c): (PEACH, ORANGE) for r in (7, 8, 9) for c in range(3)})
card(s, 8.3, 1.6, 4.43, 1.6, 'Channel reg map', 'mem0 ~ 39 設定、mem64 ~ 70 狀態；chN = 0xA009_0000 + 0x1000 × (N − 1)', tsize=17, bsize=12)
card(s, 8.3, 3.4, 4.43, 1.6, 'system_reg_map', 'sys0 ~ 16：Power Fit 模式、AGC 週期、遲滯；sys14 build_version 目前為 0', tsize=17, bsize=12)
card(s, 8.3, 5.2, 4.43, 1.6, '存取規則', 'AXI4-Lite 32-bit；沒有 reset，PS 重開機不會恢復初值；16384 = 0 dBFS', tsize=17, bsize=12)

# 15 RFDC
s = content('RFDC', 'NCO ±1.04 GHz：RF 中心 2.8 GHz')
# Zone 2 band 1.92 ~ 3.84 GHz
gx, gw, f0, f1 = 0.6, 12.13, 1.92, 3.84
fx = lambda f: gx + (f - f0) / (f1 - f0) * gw
rect(s, gx, 2.05, gw, 0.5, LT2)
tb(s, gx + 0.1, 2.13, 4.0, 0.35, 'Nyquist Zone 2 · 1.92 ~ 3.84 GHz · 3.84 GSPS', size=11, bold=True, color=DK2)
rect(s, fx(2.5), 2.75, fx(3.5) - fx(2.5), 0.4, SLATE)
tb(s, fx(2.5), 2.75, fx(3.5) - fx(2.5), 0.4, 'rev0.12：2.5 ~ 3.5 GHz（NCO ±0.84）', size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, fx(2.3), 3.3, fx(3.3) - fx(2.3), 0.45, ORANGE)
tb(s, fx(2.3), 3.3, fx(3.3) - fx(2.3), 0.45, 'rev0.1：2.3 ~ 3.3 GHz（NCO ±1.04，中心 2.8）', size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for f in (1.92, 2.3, 2.8, 3.3, 3.84):
    tb(s, fx(f) - 0.4, 1.7, 0.8, 0.3, f'{f:g}', size=10, color=DK2, align=PP_ALIGN.CENTER)
table(s, 0.6, 4.05, 8.0, [
    ['項目', 'rev0.12（同事）', 'Rex rev0.1'],
    ['NCO（兩板 ADC / DAC）', '0.84 / −0.84 GHz', '1.04 / −1.04 GHz'],
    ['ADC Outclk（z2）', '30 MHz', '240 MHz（AGC）'],
    ['z2 ADC AXIS', '未使用', '8 個 → 各通道 in_1'],
    ['z2 DAC 來源', 'hier_rx 直接接 DAC', 'cal_plus_gain_control_chN'],
    ['z2 ADC flag / clear', '未使用', '接到各通道（RXAGC）'],
], colw=[2.9, 2.55, 2.55], size=12, rowh=0.44, bold_first=True, hl={(1, 2): (PEACH, ORANGE), (2, 2): (PEACH, ORANGE)})
card(s, 8.9, 4.05, 3.83, 2.64, '注意', ['三塊板 NCO 必須一致', 'clk_adc0 也驅動 hier_mts，請同事確認 MTS 不受 240 MHz 影響', '頻率換算依設定推算，尚未量測'], tsize=17, bsize=12, bullets=True)

# 16 implementation
s = content('Implementation', '兩個 bitstream 都通過 timing')
for i, (v, l, c) in enumerate([('0.000', 'ns WNS · RFSoC_TX（0 failing endpoint）', ORANGE), ('+0.063', 'ns WNS · RFSoC_RX', GREEN),
                               ('+0.010', 'ns WHS · 兩板相同', GREEN), ('0', 'critical warning（synth / impl）', BLUE)]):
    stat(s, 0.6 + i * 3.11, 1.6, 2.81, 1.7, v, l, color=c, vsize=34)
table(s, 0.6, 3.55, 12.13, [
    ['', 'RFSoC_TX build 01', 'RFSoC_RX build 01'],
    ['LUT / FF', '213,477（50.2 %）/ 343,806（40.4 %）', '62,320（14.7 %）/ 85,553（10.1 %）'],
    ['BRAM / DSP', '566.5 tile（52.5 %）/ 2,904（68.0 %）', '224.5 tile（20.8 %）/ 0'],
    ['最差路徑', 'timer_1 → hier_rx_x0y1 debug ILA（同事設計，95 % 繞線）', 'Aurora GT 時脈 group'],
    ['輸出（output_data/01/）', 'RFSoC_TX_01.bit / .ltx / .xsa、BOOT.bin', 'RFSoC_RX_01.bit / .ltx / .xsa'],
], colw=[2.5, 5.13, 4.5], size=12, rowh=0.44, bold_first=True)
tb(s, 0.6, 5.95, 12.13, 0.8, 'Vivado 2025.2 · xczu47dr-ffvg1517-2-i。RFSoC_TX 8 個通道合計約 120K LUT、202K FF、2,904 DSP；另有 Power Fit +0.003 ns、txagc_calc +0.065 ns 的小餘裕路徑（Follow-up F1）。', size=12)

# 17 PS
s = content('PS 程式', 'RFSoC_TX 開機流程與執行期調整')
flow = [('RFDC', '初始化'), ('GPIO', 'EMIO'), ('Gain BRAM', '8 ch，讀回比對'), ('Reg map', 'system + ch1~8'), ('ADC', 'threshold × 8'),
        ('RFDC', 'reset'), ('BMC', 'handshake'), ('MTS', '+ handshake'), ('MBS', '+ handshake')]
for i, (t, sub) in enumerate(flow):
    x = 0.6 + i * 1.367
    block(s, x, 1.7, 1.2, 0.9, t, sub, style='navy' if i < 5 else 'slate', tsize=11, ssize=9)
    if i < len(flow) - 1: arrow(s, x + 1.2, 2.15, x + 1.367, 2.15)
tb(s, 0.6, 2.7, 12.13, 0.3, '前 5 步在第一次 BMC handshake 前完成；BMC 沒回 CMD 時程式會停在 handshake 等待', size=12, color=DK2)
card(s, 0.6, 3.15, 5.95, 1.5, 'BOOT.bin', 'FSBL + bit + app（沒有 PMUFW，同 Calibration_Gain_Control）；output_data/01/BOOT.bin', tsize=17, bsize=12)
card(s, 6.78, 3.15, 5.95, 1.5, 'PS 不處理', '同事的 axi_gpio_arstn / axi_delay（Aurora reset 與延遲依同事流程人工操作）', tsize=17, bsize=12)
table(s, 0.6, 4.9, 12.13, [
    ['目的（XSDB，reg_runtime/reg_rt.tcl）', '指令'],
    ['切通道 / 看狀態', 'reg_ch 3 · status all'],
    ['本板 ADC loopback / 直通', 'reg_wr signal_source_select 1 · reg_wr bypass_axis_iq 1'],
    ['數位 gain 固定 0 dB', 'reg_set gain_index_manual 0 · reg_wr gain_index_mode_sel 1'],
    ['寫 FIR 係數', 'reg_wr in_band_fir_enable 0 → reg_fir_tap <t> <I> <Q> × 16 → reg_wr in_band_fir_enable 1'],
], colw=[3.8, 8.33], size=12, rowh=0.38, bold_first=True)

# 18 verification
s = content('驗證', 'RTL 模擬全部 PASS，尚未上板')
table(s, 0.6, 1.6, 12.13, [
    ['Test', '內容', '狀態'],
    ['tb_axis_iq_source_mux', 'mem39 選擇、latency 0、tready、與 bypass 的關係（12 checks）', 'PASS'],
    ['tb_gain_index', 'reg map → cal_cdc → gain_index_mux → Gain_Compensation，300 組隨機 bit-exact（18 checks）', 'PASS'],
    ['tb_power_fit_db', 'CW / Pulse goal、2000 組隨機 pf 誤差 ≤ 1 LSB、自動鎖存（17 checks）', 'PASS'],
    ['tb_agc_emu_top', '上電、RXAGC、TX / RX 切換、bypass、類比 / 數位分配、同步、遲滯（90 checks）', 'PASS'],
    ['tb_txagc_calc_model', '對獨立行為模型，1,007,538 vectors', 'PASS'],
    ['tb_pulse_timeline', 'Power Fit → AGC → SR / Gain 整鏈時間線', 'PASS'],
    ['Delay_Cal / CORDIC', '沒有 TB；CORDIC latency 未確認', 'OPEN'],
    ['上板', 'RFSoC_TX / RFSoC_RX build 01 皆未上板', 'OPEN'],
], colw=[2.9, 7.93, 1.3], size=12, rowh=0.44, bold_first=True,
    hl={(r, 2): ((MIST, GREEN) if r <= 6 else (PEACH, RED)) for r in range(1, 9)})
tb(s, 0.6, 5.75, 12.13, 0.9, '模擬在 source 專案 Calibration_Gain_Control/verification/ 執行（Vivado 2025.2 xsim，回歸 11 個 TB，2026-10-14）；RFSoC_TX 的 RTL 與 source 逐檔相同，未在本專案重跑。多通道接線只由 BD 讀取確認。', size=12)

# 19 follow-up
s = content('Follow-up', '已知項目：下一版先處理 timing')
table(s, 0.6, 1.6, 12.13, [
    ['#', '項目', '建議', '時機'],
    ['F1', 'Timing 餘裕幾乎為 0', 'cw_power_calc_IQ 輸入加一級暫存器、txagc_calc stage 2 拆級、debug ILA probe 前加暫存器', '下一版'],
    ['F2', '無訊號時 pf = +511 且沒有 flag', '加 sticky flag（飽和、輸入 0、CW timeout）或改輸出 0 / 保持前值', '後續版本'],
    ['F3', 'Power Fit 設定 8 通道共用', '需要各通道獨立時搬到 channel reg map', '依需求'],
    ['F4', 'sys14 build_version 讀回 0', 'build_full_xsa.tcl 設定 CONFIG.BUILD_VERSION', '下一版'],
    ['F5', 'FIR 係數尚未依本板量測', '逐通道量頻響，產生 8 組 S0.15 係數', '上板後'],
], colw=[0.6, 3.2, 6.73, 1.6], size=12, rowh=0.44, bold_first=True, hl={(1, 3): (PEACH, ORANGE), (4, 3): (PEACH, ORANGE)})
card(s, 0.6, 4.4, 5.95, 2.4, 'TRM 上線後處理', ['T1 SR 腳位接出：建 port、XDC（1.8 V）、重新 build', 'T2 開機模式：依 TRM 需求決定（目前 RX mode）', 'T3 RX mode 數位 gain 與 RXAGC flag clear 的意圖確認'], dark=True, tsize=17, bsize=12, bullets=True)
card(s, 6.78, 4.4, 5.95, 2.4, '請同事確認', ['ADC Outclk 240 MHz 對 hier_mts 的影響', 'onebit_cdc_fifo 補回的設定，更新 release', 'hier_rx 的 M_AXIS_C0/C1 改成 Master'], tsize=17, bsize=12, bullets=True)

# 20 documents
s = content('文件', '文件集中在 docs/，依層級分資料夾')
cw = (12.13 - 3 * 0.3) / 4
for i, (k, t, b) in enumerate([
        ('01_System/', '系統層', ['RFSoC_TX 設計筆記', 'RFSoC_RX 設計筆記', 'VU13P 設計筆記（同事原文）']),
        ('02_RFSoC_TX/', 'RFSoC_TX', ['Top_Architecture', 'Register_Map', 'Calibration', 'Power_Fit', 'Gain_Control', 'RFDC_Configuration', 'Follow_Up', 'Reg_map / tools / history']),
        ('03_RFSoC_RX/', 'RFSoC_RX', ['Top_Architecture', 'RFDC_Configuration']),
        ('reference_rev0.12/', '同事 rev0.12', ['DESIGN', 'SIGNAL_TRANSMISSION', 'TIMING_SYNC', 'DESIGN_NOTES × 3', '原始簡報（本簡報版型）'])]):
    card(s, 0.6 + i * (cw + 0.3), 1.6, cw, 3.7, t, b, kicker=k, tsize=20, bsize=12, bullets=True)
tb(s, 0.6, 5.6, 12.13, 1.2, ['每份 md 旁有同名 docx（docs/tools/mk_all_docx.sh 產生）；索引與閱讀指引見 docs/README.md。',
                             '本簡報與產生腳本在 docs/04_Presentation/（make_ppt.py）。'], size=13, after=6)

# 21 next steps
s = content('下一步', '上板 bring-up（IF 輸出端）')
nxt = [('燒錄', 'VU13P（rev0.12）、RFSoC_RX 01 .bit / .ltx、RFSoC_TX 01 BOOT.bin（各用自己的 .ltx）'),
       ('Links', '確認 Aurora channel_up；依同事流程 release arstn（z2 → z1）後送一次 SYSREF'),
       ('暫存器', 'xsdb reg_rt.tcl status all：snapshot 計數增加、st_goal_db = −6 dBFS'),
       ('ADC loopback', 'signal_source_select = 1，確認 st_power_amp 與輸入位準一致、pf ≈ −6 − input_dBFS'),
       ('Aurora', '切回 0，Aurora 資料流動時重複，量 8 個 DAC 的 IF 輸出'),
       ('校正', '逐通道量頻響產生 FIR 係數；Delay / Phase / Gain 對齊 8 通道')]
for i, (h, d) in enumerate(nxt):
    y = 1.7 + i * 0.82
    num(s, 0.6, y, i + 1, d=0.45, color=ORANGE if i == 0 else NAVY, size=14)
    tb(s, 1.3, y + 0.05, 11.4, 0.5, [[(h + '   ', True, NAVY), (d, False, DK2)]], size=15)

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
